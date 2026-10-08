from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from utils import (
    find_data_start_row,
    is_empty,
    natural_sort_key,
    select_folder,
    trim_ghost_rows,
)


def run_data_cleaner():
    print("\n" + "="*60)
    print(" BẮT ĐẦU CHUẨN HOÁ DỮ LIỆU ĐỂ ĐƯA LÊN HỆ THỐNG")
    print("="*60)

    print("[+] Đang mở cửa sổ chọn thư mục... (Kiểm tra taskbar nếu không thấy)")
    folder_path = select_folder("Chọn thư mục chứa các file Excel cần xử lý")

    if not folder_path:
        print("[-] Đã hủy chọn thư mục. Kết thúc thao tác.")
        return

    target_dir = Path(folder_path)
    excel_files = sorted(
        [
            f
            for f in target_dir.iterdir()
            if f.is_file()
            and f.suffix in ['.xlsx', '.xls']
            and not f.name.startswith('~')
            and not f.name.startswith('DS_Tai_Lieu_So_Hoa')
        ],
        key=natural_sort_key,
    )

    if not excel_files:
        print(f"[-] Không tìm thấy file Excel nào trong thư mục {target_dir.name}")
        return

    print(f"\n[+] Đã tìm thấy {len(excel_files)} file. Bắt đầu xử lý...\n")

    yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
    red_fill = PatternFill(start_color="FFFF9999", end_color="FFFF9999", fill_type="solid")
    orange_fill = PatternFill(start_color="FFFFC000", end_color="FFFFC000", fill_type="solid")

    MANDATORY_COLS = {
        "Số hiệu bằng (B)": 2,
        "Tên văn bằng (D)": 4,
        "Ngày sinh (I)": 9,
        "Giới tính (J)": 10,
        "Dân tộc (K)": 11,
        "Quốc tịch (L)": 12,
        "Nơi sinh (M)": 13,
        "Tên trường (N)": 14,
        "Niên khoá (O)": 15,
        "Hội đồng thi (Q)": 17,
        "Địa danh cấp bằng (R)": 18,
        "Ngày cấp bằng (S)": 19,
        "Tên cơ quan (T)": 20,
        "Chức danh (U)": 21,
        "Người ký (V)": 22,
        "Trạng thái số hóa (Y)": 25,
        "Mã đơn vị (AA)": 27,
        "Tên đơn vị (AB)": 28
    }

    default_kws = ['stt', 'số hiệu bằng', 'họ và tên', 'hội đồng thi', 'nơi sinh', 'giới tính', 'dân tộc']

    files_with_issues = []
    count_processed = 0

    for file_path in excel_files:
        print(f"\n--- Đang xử lý: {file_path.name} ---")
        try:
            wb = openpyxl.load_workbook(file_path)
            target_sheet = None
            start_row = None

            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]

                found_row = find_data_start_row(
                    ws,
                    config_file="data/tu_khoa_cleaner.txt",
                    default_keywords=default_kws,
                    default_start_row=None
                )

                if found_row is not None:
                    target_sheet = ws
                    start_row = found_row
                    break

            # CHỐT CHẶN BẢO MẬT TYPE HINT CHO PYRIGHT
            if target_sheet is None or start_row is None:
                print("  [!] KHÔNG TÌM THẤY sheet nào chứa dữ liệu hợp lệ (Thiếu cột chuẩn). Bỏ qua file.")
                continue

            is_modified = False

            if target_sheet.title != "Data":
                if "Data" in wb.sheetnames:
                    wb["Data"].title = "Data_old"

                old_name = target_sheet.title
                target_sheet.title = "Data"
                is_modified = True
                print(f"  -> Đã xác định và đổi tên sheet '{old_name}' thành 'Data'.")
            else:
                print("  -> Đã xác định đúng sheet 'Data'.")

            ghost_deleted = trim_ghost_rows(target_sheet)
            if ghost_deleted > 0:
                is_modified = True
                print(f"  -> Đã gọt sạch {ghost_deleted} 'dòng ma' ở đuôi file.")

            red_rows = []
            missing_cells = []
            so_hieu_map = defaultdict(list)
            yellow_cells_count = 0

            # Lúc này start_row chắc chắn là int, hàm range() sẽ không bị cảnh báo nữa
            for r in range(start_row, target_sheet.max_row + 1):
                if all(is_empty(target_sheet.cell(row=r, column=c).value) for c in range(1, 30)):
                    continue

                if is_empty(target_sheet.cell(row=r, column=5).value):
                    for c in range(1, 30):
                        target_sheet.cell(row=r, column=c).fill = red_fill
                    red_rows.append(r)
                    is_modified = True
                else:
                    row_missing_cols = []
                    for col_name, col_idx in MANDATORY_COLS.items():
                        cell = target_sheet.cell(row=r, column=col_idx)
                        if is_empty(cell.value):
                            cell.fill = yellow_fill
                            row_missing_cols.append(col_name)
                            yellow_cells_count += 1
                            is_modified = True

                    if row_missing_cols:
                        missing_cells.append({"row": r, "cols": row_missing_cols})

                    sh_cell_val = target_sheet.cell(row=r, column=2).value
                    if not is_empty(sh_cell_val):
                        sh_str = str(sh_cell_val).strip()
                        so_hieu_map[sh_str].append(r)

            duplicate_so_hieu = {}
            for sh_str, rows in so_hieu_map.items():
                if len(rows) > 1:
                    duplicate_so_hieu[sh_str] = rows
                    is_modified = True
                    for r in rows:
                        target_sheet.cell(row=r, column=2).fill = orange_fill

            has_issue = bool(red_rows or missing_cells or duplicate_so_hieu)

            if has_issue:
                files_with_issues.append({
                    "file_name": file_path.name,
                    "red_rows": red_rows,
                    "missing_cells": missing_cells,
                    "duplicates": duplicate_so_hieu
                })
                print(f"  [!] PHÁT HIỆN LỖI: {len(red_rows)} dòng rác (Đỏ) | {yellow_cells_count} ô thiếu (Vàng) | {len(duplicate_so_hieu)} mã trùng Số hiệu bằng (Cam).")
            else:
                print("  -> OK: Dữ liệu hoàn hảo, không có lỗi.")

            if is_modified:
                wb.save(file_path)
            count_processed += 1

        except Exception as e:  # noqa: BLE001
            print(f"  [!] LỖI khi quét file: {e}")

    print("\n" + "="*60)
    print(f"[=] HOÀN TẤT CHUẨN HOÁ ({count_processed}/{len(excel_files)} file)!")

    if files_with_issues:
        print("\n" + "!"*70)
        print(" [BẢNG PHONG THẦN - CHI TIẾT CÁC FILE & DÒNG/CỘT CẦN SỬA TAY]")
        print(" -> Dòng bôi ĐỎ NHẠT : Kéo chọn và ấn Delete (dòng rác thiếu họ tên).")
        print(" -> Ô bôi VÀNG      : Bổ sung dữ liệu còn thiếu.")
        print(" -> Ô bôi CAM       : Trùng lặp 'Số hiệu bằng' (Cột B) giữa các dòng.")
        print(" -> Cột STT (A)     : Kéo chuột đánh lại sau khi xóa dòng đỏ.")
        print("!"*70)

        for issue in files_with_issues:
            print(f"\n📁 File: {issue['file_name']}")

            if issue["red_rows"]:
                rows_str = ", ".join(map(str, issue["red_rows"]))
                print(f"   ❌ DÒNG RÁC (Thiếu Họ Tên - Tô đỏ): Dòng {rows_str}")

            if issue["duplicates"]:
                print("   ⚠️  TRÙNG LẶP SỐ HIỆU BẰNG (Cột B - Tô cam):")
                for sh, rows in issue["duplicates"].items():
                    rows_str = ", ".join(f"Dòng {r}" for r in rows)
                    print(f"      + Số hiệu '{sh}': xuất hiện ở {rows_str}")

            if issue["missing_cells"]:
                print("   ⚠️  DỮ LIỆU CÒN THIẾU (Tô vàng):")
                for item in issue["missing_cells"]:
                    cols_str = ", ".join(item["cols"])
                    print(f"      + Dòng {item['row']}: Thiếu [{cols_str}]")

        print("\n" + "!"*70)
    else:
        print("\n[+] Tuyệt vời! Toàn bộ file đều hoàn hảo 100%, sẵn sàng up hệ thống.")
