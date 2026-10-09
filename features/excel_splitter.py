import json
import re
import warnings
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from thefuzz import fuzz, process

# Nhập hàm dò tìm thông minh từ utils
from utils import find_data_start_row

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Danh mục từ đồng nghĩa để nhận diện cột dữ liệu từ file nguồn
SYNONYMS_MAP = {
    "B": {
        "name": "Số hiệu bằng",
        "keywords": ["số hiệu bằng", "số hiệu", "số bằng", "so hieu", "số phôi", "so hieu bang", "số vb", "số văn bằng"]
    },
    "C": {
        "name": "Số vào sổ gốc cấp văn bằng",
        "keywords": ["số vào sổ gốc", "số vào sổ", "số sổ gốc", "số sổ", "so vao so", "so so goc", "vào sổ gốc", "số sổ cấp bằng"]
    },
    "E": {
        "name": "Họ, chữ đệm và tên",
        "keywords": ["họ và tên", "họ tên", "họ tên học sinh", "họ tên thí sinh", "họ, chữ đệm và tên", "họ và chữ đệm", "ho va ten", "họ tên người học", "tên học sinh"]
    },
    "F": {
        "name": "Mã người học",
        "keywords": ["mã người học", "mã học sinh", "ma nguoi hoc", "ma hoc sinh", "mã định danh", "mã số học sinh"]
    },
    "G": {
        "name": "Số định danh cá nhân",
        "keywords": ["số định danh cá nhân", "định danh cá nhân", "định danh", "cccd", "cmnd", "căn cước", "số cccd", "số cmnd"]
    },
    "H": {
        "name": "Hộ chiếu",
        "keywords": ["hộ chiếu", "passport", "ho chieu", "số hộ chiếu"]
    },
    "I": {
        "name": "Ngày, tháng, năm sinh",
        "keywords": ["ngày, tháng, năm sinh", "ngày tháng năm sinh", "ngày sinh", "ngay sinh", "năm sinh", "ngày tháng sinh", "sinh ngày", "ngaysinh", "ngày, tháng sinh"]
    },
    "J": {
        "name": "Giới tính",
        "keywords": ["giới tính", "phái", "nam/nữ", "nam nữ", "gioi tinh", "gt"]
    },
    "K": {
        "name": "Dân tộc",
        "keywords": ["dân tộc", "dan toc", "dt"]
    },
    "M": {
        "name": "Nơi sinh",
        "keywords": ["nơi sinh", "noi sinh", "quê quán", "nguyên quán", "tỉnh/tp nơi sinh", "tỉnh nơi sinh"]
    },
    "N": {
        "name": "Tên trường đã học",
        "keywords": ["tên trường đã học", "tên trường", "trường thpt", "trường thcs", "trường học", "trường", "đơn vị", "ten truong", "trường tốt nghiệp", "trường học sinh", "tên trường thpt"]
    },
    "P": {
        "name": "Điểm các môn thi",
        "keywords": ["điểm các môn thi", "điểm thi", "kết quả thi", "điểm xét tốt nghiệp", "điểm tn", "xếp loại tốt nghiệp"]
    },
    "Q": {
        "name": "Hội đồng thi",
        "keywords": ["hội đồng thi", "hội đồng coi thi", "hội đồng", "hđ thi", "hđ coi thi", "hội đồng thi tốt nghiệp", "hội đồng thi ghép"]
    },
    "S": {
        "name": "Ngày tháng năm cấp bằng",
        "keywords": ["ngày tháng năm cấp bằng", "ngày cấp bằng", "ngày cấp", "ngay cap", "ngày ra quyết định", "ngày ký"]
    },
    "U": {
        "name": "Chức danh người ký bằng",
        "keywords": ["chức danh người ký bằng", "chức vụ", "chức danh", "chức vụ người ký"]
    },
    "V": {
        "name": "Họ, chữ đệm, tên người ký bằng",
        "keywords": ["họ, chữ đệm, tên người ký bằng", "người ký bằng", "họ tên người ký", "người ký", "người ký quyết định"]
    },
    "W": {
        "name": "Tình trạng văn bằng",
        "keywords": ["tình trạng văn bằng", "tình trạng bằng", "xếp loại", "diện trúng tuyển"]
    },
    "X": {
        "name": "Ghi chú",
        "keywords": ["ghi chú", "ghi chu", "note", "diện ưu tiên", "chú thích"]
    },
    "AA": {
        "name": "Mã đơn vị",
        "keywords": ["mã đơn vị", "ma don vi", "mã trường", "ma truong", "mã đv", "madv", "mã cơ sở"]
    },
    "AB": {
        "name": "Tên đơn vị",
        "keywords": ["tên đơn vị", "ten don vi", "đơn vị quản lý", "tên đv", "tendv", "đơn vị", "cơ quan"]
    }
}

# Các trường thông tin tĩnh mặc định
DEFAULT_STATIC_FIELDS = {
    "D": ("Bằng tốt nghiệp phổ thông trung học", "Tên văn bằng"),
    "L": ("Việt Nam", "Quốc tịch"),
    "O": ("2022-2025", "Niên khóa"),
    "R": ("Tỉnh Vĩnh Phúc", "Địa danh nơi cơ quan cấp bằng"),
    "T": ("Sở GD&ĐT Vĩnh Phúc", "Tên cơ quan cấp bằng"),
    "Y": ("Đã số hóa đầy đủ, chính xác", "Trạng thái số hóa"),
}

def load_mapping_rules(mapping_file: str) -> dict:
    mapping = {}
    path = Path(mapping_file)
    if not path.exists():
        return mapping

    with open(mapping_file, 'r', encoding='utf-8') as f:
        for line in f:
            if '#' in line:
                line = line.split('#')[0]
            line = line.strip()
            if not line:
                continue
            if '=' in line:
                key, val = line.split('=', 1)
                key = key.strip().upper()
                val = val.strip()
                if val:
                    if (val.startswith('"') and val.endswith('"')) or (val.startswith("'") and val.endswith("'")):
                        mapping[key] = ('const', val[1:-1])
                    else:
                        mapping[key] = ('col', val.upper())
    return mapping

def save_mapping_rules(mapping_rules: dict, mapping_file: str):
    """Lưu các quy tắc ánh xạ vào file cấu hình với chú thích rõ ràng."""
    lines = [
        "# --- [Cột Template] = [Cột Demo / \"Giá trị tĩnh\"] # Tên gợi nhớ ---",
        ""
    ]
    template_col_names = {
        "A": "STT (Code đã tự động đánh số 1, 2, 3...)",
        "B": "Số hiệu bằng",
        "C": "Số vào sổ gốc cấp văn bằng",
        "D": "Tên văn bằng (Tự động điền hàng loạt)",
        "E": "Họ, chữ đệm và tên",
        "F": "Mã người học (để trống)",
        "G": "Số định danh cá nhân (để trống)",
        "H": "Hộ chiếu (để trống)",
        "I": "Ngày, tháng, năm sinh (Code tự động ép chuẩn dd/MM/yyyy)",
        "J": "Giới tính",
        "K": "Dân tộc",
        "L": "Quốc tịch",
        "M": "Nơi sinh",
        "N": "Tên trường đã học",
        "O": "Niên khóa",
        "P": "Điểm các môn thi",
        "Q": "Hội đồng thi",
        "R": "Địa danh nơi cơ quan cấp bằng đặt trụ sở",
        "S": "Ngày tháng năm cấp bằng (Code tự ép chuẩn dd/MM/yyyy)",
        "T": "Tên cơ quan cấp bằng",
        "U": "Chức danh người ký bằng",
        "V": "Họ, chữ đệm, tên người ký bằng",
        "W": "Tình trạng văn bằng",
        "X": "Ghi chú",
        "Y": "Trạng thái số hóa",
        "AA": "Mã đơn vị",
        "AB": "Tên đơn vị",
    }

    all_cols = sorted(template_col_names.keys(), key=lambda x: column_index_from_string(x))
    for col in all_cols:
        desc = template_col_names.get(col, "")
        if col == "N":
            lines.append("\n# 👇 ĐÂY LÀ MỎ NEO QUAN TRỌNG NHẤT (Bắt buộc phải có)")

        if col in mapping_rules:
            m_type, m_val = mapping_rules[col]
            if m_type == 'const':
                lines.append(f'{col} = "{m_val}" # {desc}')
            else:
                lines.append(f'{col} = {m_val} # {desc}')
        else:
            lines.append(f'{col} =   # {desc}')

    with open(mapping_file, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines) + '\n')

def extract_input_headers(ws, header_row: int) -> list:
    """
    Rút trích danh sách cột và tên tiêu đề từ file nguồn.
    Hỗ trợ cả trường hợp header 2 tầng (ghép dòng trên nếu dòng dưới rỗng).
    """
    headers = []
    max_c = min(ws.max_column, 50)
    for c in range(1, max_c + 1):
        letter = get_column_letter(c)
        val = ws.cell(row=header_row, column=c).value
        if (val is None or str(val).strip() == "") and header_row > 1:
            val = ws.cell(row=header_row - 1, column=c).value

        title = str(val).strip() if val is not None else ""
        if title:
            headers.append((letter, title))
    return headers

def auto_detect_mapping_rules(input_headers: list, existing_rules: dict | None = None) -> tuple[dict, dict]:
    """
    Tự động khớp cột từ file nguồn sang template dựa trên từ khóa đồng nghĩa và Fuzzy Matching.
    Trả về: (mapping_rules, match_info)
    """
    candidates = []
    for temp_col, meta in SYNONYMS_MAP.items():
        for in_col, in_title in input_headers:
            norm_title = in_title.strip().lower()
            score = 0
            for kw in meta["keywords"]:
                if kw == norm_title:
                    score = 100
                    break
                s = fuzz.token_set_ratio(kw, norm_title)
                if s > score:
                    score = s
            if score >= 75:
                candidates.append((score, temp_col, in_col, meta["name"], in_title))

    candidates.sort(key=lambda x: x[0], reverse=True)

    mapping_rules = {}
    match_info = {}
    assigned_temp = set()
    assigned_input = set()

    for score, temp_col, in_col, temp_name, in_title in candidates:
        if temp_col not in assigned_temp and in_col not in assigned_input:
            mapping_rules[temp_col] = ('col', in_col)
            match_info[temp_col] = (in_col, in_title, score, temp_name)
            assigned_temp.add(temp_col)
            assigned_input.add(in_col)

    # Nạp các giá trị tĩnh (ưu tiên lấy từ existing_rules nếu có, ngược lại lấy mặc định)
    for s_col, (default_val, _desc) in DEFAULT_STATIC_FIELDS.items():
        if existing_rules and s_col in existing_rules and existing_rules[s_col][0] == 'const':
            mapping_rules[s_col] = existing_rules[s_col]
        else:
            mapping_rules[s_col] = ('const', default_val)

    return mapping_rules, match_info

def print_mapping_table(mapping_rules: dict, match_info: dict, input_headers: list, split_by: str = "N"):
    """In bảng ánh xạ trực quan để người dùng kiểm tra."""
    print("\n" + "="*75)
    print("             BẢNG TỰ ĐỘNG NHẬN DIỆN VÀ ÁNH XẠ CỘT DỮ LIỆU")
    print("="*75)
    print(f" {'Cột Template':<28} <-- {'Cột Nguồn (Input)':<25} {'Độ khớp':<10}")
    print("-" * 75)

    input_map = dict(input_headers)
    dynamic_cols = [k for k, v in mapping_rules.items() if v[0] == 'col']
    dynamic_cols.sort(key=lambda x: column_index_from_string(x))

    for temp_col in dynamic_cols:
        in_col = mapping_rules[temp_col][1]
        in_title = input_map.get(in_col, "---")
        temp_name = SYNONYMS_MAP.get(temp_col, {}).get("name", f"Cột {temp_col}")
        star = " ⭐" if temp_col == split_by else ""

        score_str = ""
        if temp_col in match_info and match_info[temp_col][0] == in_col:
            score = match_info[temp_col][2]
            score_str = f"({score}%)"

        temp_label = f"[{temp_col}] {temp_name}{star}"
        src_label = f"[{in_col}] {in_title}"
        print(f" {temp_label:<28} <-- {src_label:<25} {score_str:<10}")

    print("-" * 75)
    print(" CÁC TRƯỜNG THÔNG TIN TĨNH:")
    static_cols = [k for k, v in mapping_rules.items() if v[0] == 'const']
    static_cols.sort(key=lambda x: column_index_from_string(x))
    for s_col in static_cols:
        val = mapping_rules[s_col][1]
        desc = DEFAULT_STATIC_FIELDS.get(s_col, ("", f"Cột {s_col}"))[1]
        print(f"   [{s_col}] {desc:<24} : \"{val}\"")
    print("=" * 75)

def review_and_confirm_mapping(
    mapping_rules: dict,
    match_info: dict,
    input_headers: list,
    mapping_file: str,
    split_by: str = "N"
) -> dict | None:
    """Hiển thị bảng ánh xạ và cho phép người dùng xác nhận hoặc điều chỉnh nhanh."""
    anchor_name = SYNONYMS_MAP.get(split_by, {}).get("name", f"Cột {split_by}")
    input_cols_list = [c[0] for c in input_headers]

    while True:
        print_mapping_table(mapping_rules, match_info, input_headers, split_by=split_by)

        has_anchor = (split_by in mapping_rules and mapping_rules[split_by][0] == 'col')
        if not has_anchor:
            print(f"\n [!] CẢNH BÁO QUAN TRỌNG: Chưa nhận diện được cột '{anchor_name}' (Cột {split_by})!")
            print(f"     Đây là cột BẮT BUỘC để phân tách dữ liệu theo {anchor_name}.")

        print("\n Bạn muốn làm gì?")
        if has_anchor:
            print(" [Enter] Đồng ý toàn bộ và BẮT ĐẦU TÁCH FILE ngay")
        else:
            print(f" [Enter] Chọn ngay cột '{anchor_name}' (Cột {split_by})")
        print(" [1]     Sửa / gán lại một cột động (vd: đổi cột Họ tên, Số hiệu, Mã đơn vị...)")
        print(" [2]     Chỉnh sửa thông tin tĩnh (Niên khóa, Tên bằng, Địa danh...)")
        print(" [3]     Dùng lại nguyên văn quy tắc cũ từ file 'quy_tac_anh_xa.txt'")
        print(" [0]     Hủy thao tác")

        choice = input("\n -> Lựa chọn của bạn [Mặc định: Enter]: ").strip()

        if choice == "":
            if not has_anchor:
                print("\n--- DANH SÁCH CÁC CỘT FILE NGUỒN HIỆN CÓ ---")
                for col_l, col_t in input_headers:
                    print(f"   [{col_l}] {col_t}")
                col_choice = input(f"\nNhập chữ cái cột chứa {anchor_name} (vd: G): ").strip().upper()
                if col_choice in input_cols_list:
                    mapping_rules[split_by] = ('col', col_choice)
                    match_info[split_by] = (col_choice, dict(input_headers).get(col_choice, ""), 100, anchor_name)
                    print(f" -> Đã gán thành công: Cột {split_by} <-- [{col_choice}].")
                else:
                    print(" [!] Cột vừa nhập không có trong file nguồn. Vui lòng thử lại.")
                continue

            save_mapping_rules(mapping_rules, mapping_file)
            print(f"\n[+] Đã tự động cập nhật quy tắc vào file: {Path(mapping_file).name}")
            return mapping_rules

        elif choice == "0":
            return None

        elif choice == "1":
            print("\n--- DANH SÁCH CÁC CỘT FILE NGUỒN HIỆN CÓ ---")
            for col_l, col_t in input_headers:
                print(f"   [{col_l}] {col_t}")
            t_col = input("\nNhập chữ cái cột Template muốn sửa (vd: B, E, N, Q, AA, AB... hoặc 'x' để quay lại): ").strip().upper()
            if t_col in ['X', '']:
                continue

            s_col = input(f"Nhập chữ cái cột Nguồn muốn gán cho [{t_col}] (hoặc gõ '0' để xóa ánh xạ cột này): ").strip().upper()
            if s_col == '0':
                if t_col in mapping_rules:
                    del mapping_rules[t_col]
                if t_col in match_info:
                    del match_info[t_col]
                print(f" -> Đã xóa ánh xạ của cột [{t_col}].")
            elif s_col in input_cols_list:
                mapping_rules[t_col] = ('col', s_col)
                in_name = dict(input_headers).get(s_col, "")
                temp_name = SYNONYMS_MAP.get(t_col, {}).get("name", f"Cột {t_col}")
                match_info[t_col] = (s_col, in_name, 100, temp_name)
                print(f" -> Đã gán: [{t_col}] <-- [{s_col}] ({in_name}).")
            else:
                print(f" [!] Cột nguồn '{s_col}' không tồn tại trong danh sách cột file nguồn.")

        elif choice == "2":
            print("\n--- CHỈNH SỬA THÔNG TIN TĨNH (Nhấn [Enter] để giữ nguyên giá trị cũ) ---")
            for s_col in ["O", "D", "T", "R", "L", "Y"]:
                if s_col in DEFAULT_STATIC_FIELDS:
                    desc = DEFAULT_STATIC_FIELDS[s_col][1]
                    cur_val = mapping_rules.get(s_col, ('const', DEFAULT_STATIC_FIELDS[s_col][0]))[1]
                    new_val = input(f" - [{s_col}] {desc} [{cur_val}]: ").strip()
                    if new_val:
                        mapping_rules[s_col] = ('const', new_val)
            print(" -> Đã cập nhật xong thông tin tĩnh!")

        elif choice == "3":
            if Path(mapping_file).exists():
                mapping_rules = load_mapping_rules(mapping_file)
                match_info = {}
                print(f" -> Đã tải lại toàn bộ quy tắc từ {Path(mapping_file).name}!")
            else:
                print(" [!] File quy tắc cũ không tồn tại.")

        else:
            print(" [!] Lựa chọn không hợp lệ, vui lòng thử lại.")

def load_school_catalog(csv_file: str) -> dict:
    catalog = {}
    with open(csv_file, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or ';' not in line or line.startswith('ma_truong'):
                continue
            ma_truong, ten_truong = line.split(';', 1)
            catalog[ten_truong.strip()] = ma_truong.strip()
    return catalog

def run_excel_splitter(input_file: str, template_file: str, mapping_file: str, output_dir: str):
    # 1. KIỂM TRA TÍNH HỢP LỆ CỦA FILE TEMPLATE TRƯỚC KHI CHẠY
    try:
        wb_temp = openpyxl.load_workbook(template_file)
        if 'Data' not in wb_temp.sheetnames:
            print("[-] LỖI: File template không có sheet 'Data'. Vui lòng kiểm tra lại file mẫu!")
            return
        wb_temp.close()
    except Exception as e:  # noqa: BLE001
        print(f"[-] LỖI: Không thể mở file template {template_file}. Chi tiết: {e}")
        return

    # 2. CHỌN TIÊU CHÍ TÁCH FILE
    print("\n" + "="*60)
    print(" BẮT ĐẦU TÁCH FILE DỮ LIỆU EXCEL")
    print("="*60)
    print(" Bạn muốn tách file theo tiêu chí nào?")
    print("   1. Tách theo Tên trường đã học (Cột N) [Mặc định]")
    print("   2. Tách theo Hội đồng thi (Cột Q)")
    split_choice = input(" -> Lựa chọn (1/2) [Mặc định: 1]: ").strip()

    if split_choice == "2":
        split_by = "Q"
        split_name = "Hội đồng thi"
        prefix = "HDT_"
    else:
        split_by = "N"
        split_name = "Tên trường đã học"
        prefix = ""

    print(f"\n[+] Đã chọn tiêu chí: Tách theo '{split_name}' (Cột {split_by})")

    # 3. DÙNG OPENPYXL ĐỂ DÒ TÌM DÒNG BẮT ĐẦU VÀ RÚT TRÍCH TIÊU ĐỀ FILE INPUT
    try:
        wb_in = openpyxl.load_workbook(input_file, data_only=True)
        ws_in = wb_in.active

        default_kws = ['stt', 'số hiệu bằng', 'họ và tên', 'hội đồng thi', 'nơi sinh', 'giới tính', 'dân tộc']
        start_row = find_data_start_row(
            ws_in,
            config_file="data/tu_khoa_splitter.txt",
            default_keywords=default_kws,
            default_start_row=None
        )

        if start_row is None:
            wb_in.close()
            print("[-] LỖI: File input không chứa cấu trúc bảng dữ liệu hợp lệ (Không đủ từ khóa).")
            return

        print(f"  -> Đã nhận diện cấu trúc file nguồn. Dữ liệu bắt đầu từ dòng: {start_row}")

        # Tự động rút trích danh sách cột tiêu đề từ dòng ngay trước start_row
        header_row = start_row - 1
        input_headers = extract_input_headers(ws_in, header_row)
        wb_in.close()

    except Exception as e:  # noqa: BLE001
        print(f"[-] LỖI đọc file input bằng openpyxl: {e}")
        return

    # 4. TỰ ĐỘNG KHỚP CỘT THÔNG MINH VÀ CHO NGƯỜI DÙNG XÁC NHẬN
    existing_rules = load_mapping_rules(mapping_file) if Path(mapping_file).exists() else {}
    auto_rules, match_info = auto_detect_mapping_rules(input_headers, existing_rules)
    confirmed_rules = review_and_confirm_mapping(
        auto_rules,
        match_info,
        input_headers,
        mapping_file,
        split_by=split_by
    )

    if not confirmed_rules:
        print("[-] Đã hủy thao tác tách file.")
        return

    mapping_rules = confirmed_rules

    # 5. ĐỌC DỮ LIỆU BẰNG PANDAS VÀ TIẾN HÀNH TÁCH FILE
    try:
        df = pd.read_excel(input_file, sheet_name=0, header=None)
    except Exception as e:  # noqa: BLE001
        print(f"[-] LỖI đọc file input bằng Pandas: {e}")
        return

    if split_by not in mapping_rules or mapping_rules[split_by][0] != 'col':
        print(f"[-] LỖI: Bắt buộc cấu hình ánh xạ cho cột '{split_name}' (Cột {split_by}) để tách file.")
        return

    _map_type, split_col_letter = mapping_rules[split_by]
    split_col_idx = column_index_from_string(split_col_letter) - 1

    # Cắt DataFrame từ dòng start_row (index trong Pandas là start_row - 1)
    df_data = df.iloc[start_row - 1:].copy()
    df_data = df_data.dropna(subset=[split_col_idx])

    raw_groups = df_data[split_col_idx].unique()
    print(f"\n[+] Đã tìm thấy {len(raw_groups)} nhóm {split_name} cần tách.")
    print("\n[+] Bắt đầu tách và kết xuất Excel...\n")

    for raw_group in raw_groups:
        raw_group_str = str(raw_group).strip()
        if not raw_group_str:
            continue

        df_group = df_data[df_data[split_col_idx] == raw_group]

        wb = openpyxl.load_workbook(template_file)
        ws = wb['Data']

        current_out_row = 3
        stt_counter = 1

        for _, row in df_group.iterrows():
            ws.cell(row=current_out_row, column=1, value=stt_counter)

            for temp_letter, (m_type, m_val) in mapping_rules.items():
                if temp_letter == "A":
                    continue
                col_idx = column_index_from_string(temp_letter)
                if m_type == 'const':
                    ws.cell(row=current_out_row, column=col_idx, value=m_val)
                elif m_type == 'col':
                    demo_idx = column_index_from_string(m_val) - 1
                    if demo_idx < len(row):
                        val = row.iloc[demo_idx]
                        if pd.isna(val):
                            val = ""
                        else:
                            if temp_letter in ["I", "S"]:
                                try:
                                    val = pd.to_datetime(val, dayfirst=True).strftime("%d/%m/%Y")
                                except Exception:  # noqa: BLE001
                                    val = str(val).strip()
                            elif isinstance(val, pd.Timestamp):
                                val = val.strftime("%d/%m/%Y")
                            elif isinstance(val, float) and val.is_integer():
                                val = str(int(val))
                            else:
                                val = str(val).strip()

                        ws.cell(row=current_out_row, column=col_idx, value=val)

            current_out_row += 1
            stt_counter += 1

        safe_group_name = re.sub(r'[<>:"/\\|?*]', '_', raw_group_str).strip()
        num_students = len(df_group)

        file_name = f"{prefix}{safe_group_name}_{num_students}.xlsx"

        output_path = Path(output_dir) / file_name
        wb.save(output_path)
        print(f"  -> Đã tạo: {file_name}")

    print(f"\n[=] HOÀN TẤT! Toàn bộ file đã được lưu tại: {output_dir}\n")
