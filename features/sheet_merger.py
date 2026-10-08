from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

from utils import (
    find_data_start_row,
    is_empty,
    open_file_in_os,
    save_file,
    select_file,
    trim_ghost_rows,
)

PASTEL_COLORS = [
    "FFFFFF", "F4F6F6", "EBF5FB", "FEF9E7", "EAFAF1", "F5EEF8"
]

def copy_cell(source_cell, target_cell, fill_override=None):
    target_cell.value = source_cell.value
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = copy(source_cell.number_format)

    if fill_override:
        target_cell.fill = fill_override
    elif source_cell.has_style and source_cell.fill:
        target_cell.fill = copy(source_cell.fill)

def extract_header(ws, header_row: int) -> list:
    """Rút trích dòng tiêu đề của bảng, loại bỏ các ô trống ở cuối để so sánh chuẩn xác."""
    if header_row < 1:
        return []

    header = []
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        clean_val = " ".join(str(val).split()).lower() if val is not None else ""
        header.append(clean_val)

    while header and header[-1] == "":
        header.pop()

    return header

def print_header_diff(base_header: list, current_header: list):
    """Hiển thị chi tiết sai lệch giữa khuôn mẫu chuẩn và sheet hiện tại."""
    print("\n      [!] PHÁT HIỆN LỆCH CẤU TRÚC CỘT!")
    print(f"          - Chuẩn (Sheet 1) có {len(base_header)} cột: {base_header}")
    print(f"          - Sheet này đang có {len(current_header)} cột: {current_header}")

    base_set = set(base_header)
    curr_set = set(current_header)

    missing = [c for c in base_header if c not in curr_set]
    extra = [c for c in current_header if c not in base_set]

    if missing:
        print(f"          -> Cột THIẾU so với chuẩn: {missing}")
    if extra:
        print(f"          -> Cột THỪA so với chuẩn: {extra}")

    diff_positions = []
    max_len = max(len(base_header), len(current_header))
    for idx in range(max_len):
        b_val = base_header[idx] if idx < len(base_header) else "(không có)"
        c_val = current_header[idx] if idx < len(current_header) else "(không có)"
        if b_val != c_val:
            diff_positions.append(f"Cột {idx + 1} [Chuẩn: '{b_val}' != Sheet này: '{c_val}']")

    if diff_positions:
        displayed = diff_positions[:5]
        suffix = f" ... (+{len(diff_positions) - 5} cột khác)" if len(diff_positions) > 5 else ""
        print(f"          -> Vị trí sai lệch: {'; '.join(displayed)}{suffix}")

def run_sheet_merger():
    print("\n" + "="*60)
    print(" BẮT ĐẦU GỘP CÁC SHEET (CÓ KIỂM TRA KHUÔN MẪU)")
    print("="*60)

    print("[+] Đang mở cửa sổ chọn file nguồn... (Kiểm tra taskbar nếu không thấy)")
    input_file = select_file("Chọn File Excel chứa nhiều sheet cần gộp")
    if not input_file:
        print("[-] Đã hủy chọn file.")
        return

    print("\n[+] Đang mở cửa sổ chọn nơi lưu file gộp...")
    default_out_name = Path(input_file).stem + "_Gop.xlsx"
    output_file = save_file(title="Lưu file Excel đã gộp", default_name=default_out_name)
    if not output_file:
        print("[-] Đã hủy lưu file.")
        return

    print(f"\n--- Đang xử lý file: {Path(input_file).name} ---")

    try:
        wb_in = openpyxl.load_workbook(input_file)
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active

        if ws_out is None:
            print("[-] LỖI: Không thể khởi tạo Worksheet cho file gộp.")
            return

        ws_out.title = "Data_Gop"

        current_out_row = 1
        is_first_sheet = True
        base_header = []

        default_kws = ['stt', 'số hiệu bằng', 'họ và tên', 'hội đồng thi', 'nơi sinh', 'giới tính', 'dân tộc']
        valid_sheet_count = 0

        # Đã xóa enumerate và sheet_idx dư thừa
        for sheet_name in list(wb_in.sheetnames):
            if sheet_name not in wb_in.sheetnames:
                print(f"  -> Sheet '{sheet_name}' không còn tồn tại trong file (Bỏ qua).")
                continue

            ws_in = wb_in[sheet_name]
            print(f"  -> Đang đọc sheet: '{sheet_name}'...", end="")

            start_row = find_data_start_row(
                ws_in,
                config_file="data/tu_khoa_merger.txt",
                default_keywords=default_kws,
                default_start_row=None
            )

            if start_row is None:
                print(" (Bỏ qua: Không phải sheet dữ liệu chuẩn)")
                continue

            trim_ghost_rows(ws_in)
            header_row = start_row - 1
            current_header = extract_header(ws_in, header_row)

            color_hex = PASTEL_COLORS[valid_sheet_count % len(PASTEL_COLORS)]
            sheet_fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")

            if is_first_sheet:
                base_header = current_header
                print(" (Đã lấy làm chuẩn khuôn mẫu)")

                for r in range(1, start_row):
                    for c in range(1, ws_in.max_column + 1):
                        src_cell = ws_in.cell(row=r, column=c)
                        tgt_cell = ws_out.cell(row=current_out_row, column=c)
                        copy_cell(src_cell, tgt_cell)
                    current_out_row += 1

                for col_letter, col_dim in ws_in.column_dimensions.items():
                    ws_out.column_dimensions[col_letter].width = col_dim.width

                is_first_sheet = False
            else:
                if current_header != base_header:
                    skip_sheet = False

                    while current_header != base_header:
                        print_header_diff(base_header, current_header)

                        print("\n      Bạn muốn xử lý thế nào?")
                        print("      1. Mở file Excel để sửa -> Lưu file -> Quét lại sheet này")
                        print("      2. Bỏ qua (Skip) sheet này, tiếp tục gộp các sheet khác")
                        print("      3. Cố tình Ép gộp (Force merge) - Khuyên không nên dùng")
                        print("      0. Hủy toàn bộ quá trình gộp để đi sửa file")

                        choice = input("      -> Lựa chọn (1/2/3/0) [Mặc định: 1]: ").strip()
                        if choice == "":
                            choice = "1"

                        if choice == '0':
                            print("[-] Đã hủy tiến trình gộp file.")
                            return
                        elif choice == '2':
                            print(f"      -> Đã bỏ qua sheet '{sheet_name}'.\n")
                            skip_sheet = True
                            break
                        elif choice == '3':
                            print("      -> Cảnh báo: Dữ liệu đang được ép gộp!")
                            break
                        elif choice == '1':
                            print(f"\n      [+] Đang mở file: {Path(input_file).name}...")
                            open_file_in_os(input_file)
                            print(f"      [i] Hướng dẫn: Chuyển đến sheet '{sheet_name}', chỉnh sửa cấu trúc cột cho khớp chuẩn, sau đó bấm LƯU FILE (Save / Ctrl+S / Cmd+S).")

                            rescanned_ok = False
                            while True:
                                sub = input("\n      -> Nhấn [Enter] sau khi đã LƯU file để quét lại (hoặc gõ 'b' để quay lại menu): ").strip()
                                if sub.lower() in ['b', 'back']:
                                    break

                                print(f"      [*] Đang tải lại file và quét sheet '{sheet_name}'...")
                                try:
                                    wb_in = openpyxl.load_workbook(input_file)
                                except Exception as e:
                                    print(f"      [!] Lỗi khi đọc file (có thể file đang được Excel ghi/khóa): {e}")
                                    print("          Vui lòng kiểm tra đã lưu file xong trong Excel và thử lại.")
                                    continue

                                if sheet_name not in wb_in.sheetnames:
                                    print(f"      [!] Không tìm thấy sheet '{sheet_name}' trong file sau khi tải lại!")
                                    print(f"          Danh sách sheet hiện có: {wb_in.sheetnames}")
                                    print(f"          Vui lòng kiểm tra lại tên sheet '{sheet_name}' và lưu file.")
                                    continue

                                ws_in = wb_in[sheet_name]
                                new_start_row = find_data_start_row(
                                    ws_in,
                                    config_file="data/tu_khoa_merger.txt",
                                    default_keywords=default_kws,
                                    default_start_row=None
                                )

                                if new_start_row is None:
                                    print("      [!] Cảnh báo: Không nhận diện được dòng tiêu đề dữ liệu hợp lệ sau khi sửa!")
                                    print(f"          Tiêu đề cần có các từ khóa như: {', '.join(default_kws[:4])}...")
                                    continue

                                trim_ghost_rows(ws_in)
                                header_row = new_start_row - 1
                                current_header = extract_header(ws_in, header_row)
                                start_row = new_start_row

                                if current_header == base_header:
                                    print("      [✓] Tuyệt vời! Cấu trúc cột đã khớp 100% với sheet chuẩn.")
                                    rescanned_ok = True
                                    break
                                else:
                                    print("      [!] Đã quét lại nhưng cấu trúc cột vẫn chưa khớp:")
                                    print_header_diff(base_header, current_header)

                            if rescanned_ok:
                                break
                        else:
                            print("      [!] Nhập sai, vui lòng chọn 1, 2, 3 hoặc 0.")

                    if skip_sheet:
                        continue
                else:
                    print(" (Khuôn mẫu khớp 100%)")

            data_count = 0
            for r in range(start_row, ws_in.max_row + 1):
                if all(is_empty(ws_in.cell(row=r, column=c).value) for c in range(1, 31)):
                    continue

                for c in range(1, ws_in.max_column + 1):
                    src_cell = ws_in.cell(row=r, column=c)
                    tgt_cell = ws_out.cell(row=current_out_row, column=c)
                    copy_cell(src_cell, tgt_cell, fill_override=sheet_fill)

                current_out_row += 1
                data_count += 1

            valid_sheet_count += 1
            print(f"     + Đã nối {data_count} dòng dữ liệu.")

        wb_out.save(output_file)
        print("\n" + "="*60)
        print(f"[=] HOÀN TẤT! Đã gộp thành công {valid_sheet_count} sheet hợp lệ.")
        print(f"[=] File kết quả lưu tại: {output_file}")

    except Exception as e:  # noqa: BLE001
        print(f"  [!] LỖI TỔNG THỂ: {e}")
